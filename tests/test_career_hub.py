"""
tests/test_career_hub.py
=========================
Target role setting, resume upload, and skill-gap analysis.
"""

import io
import re
import pytest


def get_csrf_from(resp):
    match = re.search(r'name="csrf-token" content="([^"]+)"', resp.get_data(as_text=True))
    return match.group(1) if match else None


def make_test_docx_bytes():
    """Builds a minimal, genuinely valid .docx resume in memory using
    python-docx (already a project dependency) — a real file, not a stub,
    so it exercises resume_parser.py's actual extraction code."""
    from docx import Document

    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("Skills")
    doc.add_paragraph("Python, SQL, Machine Learning, Communication")
    doc.add_paragraph("Projects")
    doc.add_paragraph("Built a churn-prediction model using scikit-learn and deployed it with Flask.")
    doc.add_paragraph("Education")
    doc.add_paragraph("B.Sc. Computer Science")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


class TestSetTargetRole:
    def test_set_valid_target_role(self, auth_client, app_module, test_user):
        token = get_csrf_from(auth_client.get("/career"))
        resp = auth_client.post(
            "/career/set_target",
            data={
                "target_role": "Data Scientist",
                "csrf_token": token,
            },
            follow_redirects=True,
        )
        assert resp.status_code == 200
        assert app_module.db.get_user_target_role(str(test_user["id"])) == "Data Scientist"

    def test_set_invalid_target_role_rejected(self, auth_client, app_module, test_user):
        token = get_csrf_from(auth_client.get("/career"))
        auth_client.post(
            "/career/set_target",
            data={
                "target_role": "Astronaut (Not A Real Option)",
                "csrf_token": token,
            },
            follow_redirects=True,
        )
        assert app_module.db.get_user_target_role(str(test_user["id"])) is None

    def test_set_target_role_requires_csrf(self, auth_client, app_module, test_user):
        auth_client.post("/career/set_target", data={"target_role": "Data Scientist"}, follow_redirects=True)
        assert app_module.db.get_user_target_role(str(test_user["id"])) is None


class TestResumeUpload:
    def test_upload_valid_docx_resume(self, auth_client, app_module, test_user):
        token = get_csrf_from(auth_client.get("/career"))
        data = {
            "resume": (io.BytesIO(make_test_docx_bytes()), "resume.docx"),
            "csrf_token": token,
        }
        resp = auth_client.post(
            "/career/upload_resume", data=data, content_type="multipart/form-data", follow_redirects=True
        )
        assert resp.status_code == 200
        stored = app_module.db.get_latest_resume(str(test_user["id"]))
        assert stored is not None
        assert "Python" in stored["candidate_skills"] or any("Python" in s for s in stored["candidate_skills"])

    def test_upload_rejects_disallowed_extension(self, auth_client):
        token = get_csrf_from(auth_client.get("/career"))
        data = {
            "resume": (io.BytesIO(b"not a real resume"), "resume.exe"),
            "csrf_token": token,
        }
        resp = auth_client.post(
            "/career/upload_resume", data=data, content_type="multipart/form-data", follow_redirects=True
        )
        # Should not 500 — the parser raises a user-facing ResumeParseError (422)
        assert resp.status_code in (200, 400, 422)


class TestCareerAnalysis:
    def test_analysis_without_target_role_or_resume(self, auth_client):
        """A brand-new user with nothing set yet should get a graceful
        response, not a 500."""
        resp = auth_client.get("/career/analysis")
        assert resp.status_code in (200, 400)

    def test_analysis_with_target_role_set(self, auth_client, app_module, test_user):
        token = get_csrf_from(auth_client.get("/career"))
        auth_client.post("/career/set_target", data={"target_role": "Data Scientist", "csrf_token": token})
        resp = auth_client.get("/career/analysis")
        assert resp.status_code == 200
        assert resp.is_json
